"""Tests for the résumé renderer.

Covers:
- DOCX is written to disk and is a valid docx package
- Name, contact line, summary, experience bullets, skills, and education
  all appear in the rendered document
- Graceful behavior when sections are missing (empty education, no skills)
- convert_to_pdf returns None when docx2pdf/Word is unavailable — it must
  never raise and crash the pipeline
"""

from __future__ import annotations

from pathlib import Path

import pytest

from career_ai.resume.renderer import convert_to_pdf, render_docx


@pytest.fixture
def full_resume() -> dict:
    return {
        "name": "Jane Doe",
        "contact": {
            "email": "jane@example.com",
            "phone": "+1-555-0100",
            "location": "Remote",
            "links": ["https://linkedin.com/in/janedoe", "https://github.com/janedoe"],
        },
        "summary": "Senior Backend Engineer with 7 years scaling Python services.",
        "experience": [
            {
                "company": "Acme Corp",
                "title": "Staff Engineer",
                "dates": "2021 - Present",
                "location": "Remote",
                "bullets": [
                    "Led migration of monolith to 12 microservices serving 2M RPS",
                    "Mentored 5 engineers; reduced on-call pages by 60%",
                ],
            },
            {
                "company": "Widgets Inc",
                "title": "Senior Engineer",
                "dates": "2018 - 2021",
                "location": "",
                "bullets": ["Built payment reconciliation system processing $40M/mo"],
            },
        ],
        "skills": ["Python", "PostgreSQL", "Kubernetes", "AWS"],
        "education": [
            {"school": "MIT", "degree": "B.S. Computer Science", "dates": "2014 - 2018"},
        ],
    }


def test_render_docx_writes_file(full_resume, tmp_path: Path) -> None:
    out = tmp_path / "resume.docx"
    returned = render_docx(full_resume, out)
    assert returned == out
    assert out.exists()
    assert out.stat().st_size > 1000  # a real docx is a zip with multiple parts


def test_rendered_docx_contains_all_content(full_resume, tmp_path: Path) -> None:
    out = tmp_path / "resume.docx"
    render_docx(full_resume, out)

    from docx import Document

    doc = Document(str(out))
    text = "\n".join(p.text for p in doc.paragraphs)

    # Header
    assert "Jane Doe" in text
    assert "jane@example.com" in text
    assert "linkedin.com/in/janedoe" in text

    # Section headings (upper-case per renderer convention)
    assert "SUMMARY" in text
    assert "EXPERIENCE" in text
    assert "SKILLS" in text
    assert "EDUCATION" in text

    # Content
    assert "Staff Engineer" in text
    assert "Acme Corp" in text
    assert "2M RPS" in text
    assert "60%" in text
    assert "Python" in text
    assert "MIT" in text
    assert "B.S. Computer Science" in text


def test_render_docx_handles_missing_sections(tmp_path: Path) -> None:
    minimal = {
        "name": "John Smith",
        "contact": {"email": "j@s.co", "phone": "", "location": "", "links": []},
        "summary": "",
        "experience": [],
        "skills": [],
        "education": [],
    }
    out = tmp_path / "minimal.docx"
    render_docx(minimal, out)
    assert out.exists()

    from docx import Document

    doc = Document(str(out))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "John Smith" in text
    # No section headings because every list is empty
    assert "EXPERIENCE" not in text
    assert "EDUCATION" not in text


def test_render_docx_creates_parent_dir(full_resume, tmp_path: Path) -> None:
    nested = tmp_path / "a" / "b" / "c" / "resume.docx"
    render_docx(full_resume, nested)
    assert nested.exists()


def test_render_docx_uses_list_bullet_style(full_resume, tmp_path: Path) -> None:
    """Bullets must use Word's built-in 'List Bullet' style so ATS parsers see them."""
    out = tmp_path / "resume.docx"
    render_docx(full_resume, out)

    from docx import Document

    doc = Document(str(out))
    bullet_paragraphs = [p for p in doc.paragraphs if p.style.name == "List Bullet"]
    # Three bullets total across the two roles
    assert len(bullet_paragraphs) == 3


def test_convert_to_pdf_returns_none_when_conversion_unavailable(
    full_resume, tmp_path: Path, monkeypatch
) -> None:
    """docx2pdf may not be installed or Word/LibreOffice may be missing —
    convert_to_pdf must return None rather than raise."""
    out_docx = tmp_path / "resume.docx"
    out_pdf = tmp_path / "resume.pdf"
    render_docx(full_resume, out_docx)

    # Simulate docx2pdf import failure
    import sys
    monkeypatch.setitem(sys.modules, "docx2pdf", None)

    result = convert_to_pdf(out_docx, out_pdf)
    assert result is None
    assert not out_pdf.exists()
